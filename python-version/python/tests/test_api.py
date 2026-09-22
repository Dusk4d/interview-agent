from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor
from dataclasses import replace
from io import BytesIO
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from interview_agent.api import create_app
from interview_agent.config import Settings

RESUME = """
张三
电话：13812345678 邮箱：demo@example.com

专业技能
Java、Python、Spring Boot、Redis、MySQL、Kafka、Docker

项目经历
星轨推荐引擎 2024.01-2024.08
负责推荐服务的缓存与消息链路，使用 Redis 和 Kafka，补充监控与压测验证。
将故障恢复流程整理为自动化脚本，和团队一起完成上线。

教育经历
示例大学 软件工程 2020.09-2024.06
"""


def client(tmp_path: Path) -> TestClient:
    settings = replace(
        Settings.load(), storage_mode="memory", data_dir=tmp_path, llm_provider="mock",
        max_questions=3, max_question_skips=2,
    )
    return TestClient(create_app(settings))


def import_resume(c: TestClient) -> dict:
    response = c.post("/api/resumes/text", json={"text": RESUME, "fileName": "resume.txt"})
    assert response.status_code == 200, response.text
    return response.json()


def test_health_static_and_resume_privacy(tmp_path: Path):
    c = client(tmp_path)
    health = c.get("/api/health").json()
    assert health["knowledgeItems"] >= 20
    assert health["retrievalTopK"] == 5 and health["retrievalMinScore"] == 0.08
    assert c.get("/index.html").status_code == 200
    imported = import_resume(c)
    assert imported["factCount"] >= 3
    detail = c.get(f"/api/resumes/{imported['resumeId']}").json()
    assert "13812345678" not in detail["text"]
    assert "demo@example.com" not in detail["text"]
    assert "[手机号已隐藏]" in detail["text"]
    assert detail["projects"]
    assert "Redis" in detail["techStack"]


def test_project_interview_full_loop_and_report(tmp_path: Path):
    c = client(tmp_path)
    resume_id = import_resume(c)["resumeId"]
    session = c.post("/api/interviews", json={"resumeId": resume_id, "mode": "PROJECT", "maxQuestions": 1}).json()
    question = c.post(f"/api/interviews/{session['id']}/next-question").json()
    assert question["sourceIds"] and question["citations"]
    result = c.post(f"/api/interviews/{session['id']}/answers", json={
        "questionId": question["id"],
        "content": "首先我负责 Redis 缓存与 Kafka 消息链路。其次通过压测和监控验证延迟与异常恢复，最后复盘边界条件。",
    })
    assert result.status_code == 200, result.text
    body = result.json()
    assert len(body["evaluation"]["dimensions"]) == 4
    assert body["evaluation"]["degraded"] is True
    assert body["nextAction"] == "FINISH"
    assert body["session"]["status"] == "NEXT_QUESTION"
    report = c.post(f"/api/interviews/{session['id']}/report")
    assert report.status_code == 200, report.text
    assert report.json()["report"]["answerCount"] == 1
    assert "我的回答" in report.json()["markdown"]
    assert c.get(f"/api/interviews/{session['id']}/report.md").headers["content-type"].startswith("text/markdown")


def test_question_planner_tracks_project_and_avoids_covered_topic_repeats(tmp_path: Path):
    c = client(tmp_path)
    two_project_resume = RESUME.replace(
        "\n教育经历",
        "\n云帆监控平台 2023.03-2023.12\n负责告警聚合与可观测性链路，使用 MySQL 并编写故障演练脚本。\n\n教育经历",
    )
    imported = c.post(
        "/api/resumes/text", json={"text": two_project_resume, "fileName": "resume.txt"},
    )
    assert imported.status_code == 200
    resume_id = imported.json()["resumeId"]
    project_labels = [item["name"] for item in c.get(f"/api/resumes/{resume_id}").json()["projects"]]
    assert len(project_labels) == 2
    session_id = c.post(
        "/api/interviews", json={"resumeId": resume_id, "mode": "PROJECT", "maxQuestions": 2},
    ).json()["id"]

    first = c.post(f"/api/interviews/{session_id}/next-question").json()
    session = c.get(f"/api/interviews/{session_id}").json()
    stored = c.app.state.repo.get("sessions", session_id)
    assert project_labels[0] in first["content"]
    assert session["coveredTopics"] == [project_labels[0]]
    assert stored.activeProject == project_labels[0]

    answered = c.post(f"/api/interviews/{session_id}/answers", json={
        "questionId": first["id"],
        "content": "首先说明个人职责，然后解释缓存机制、边界条件和测试验证结果。",
    })
    assert answered.status_code == 200
    second = c.post(f"/api/interviews/{session_id}/next-question").json()
    assert second["content"] != first["content"]
    assert second["type"] != first["type"]
    assert project_labels[1] in second["content"]
    assert c.get(f"/api/interviews/{session_id}").json()["coveredTopics"] == project_labels
    assert c.app.state.repo.get("sessions", session_id).activeProject == project_labels[1]


def test_report_radar_uses_ability_domains_not_single_answer_rubric(tmp_path: Path):
    c = client(tmp_path)
    resume_id = import_resume(c)["resumeId"]
    session_id = c.post(
        "/api/interviews",
        json={"resumeId": resume_id, "mode": "PROJECT", "maxQuestions": 2},
    ).json()["id"]
    question = c.post(f"/api/interviews/{session_id}/next-question").json()
    answer = c.post(
        f"/api/interviews/{session_id}/answers",
        json={"questionId": question["id"], "content": "首先说明职责，然后解释缓存机制与测试验证。"},
    ).json()
    evaluation_id = answer["evaluation"]["id"]
    evaluation = c.app.state.repo.get("evaluations", evaluation_id)
    c.app.state.repo.save("evaluations", evaluation.model_copy(update={"dimensions": evaluation.dimensions[:1]}))

    report = c.post(f"/api/interviews/{session_id}/report").json()["report"]
    radar = {item["key"]: item for item in report["abilityRadar"]}
    assert radar["project_expression"]["sample"] == 1
    assert radar["project_expression"]["score"] == answer["evaluation"]["totalScore"]
    assert radar["middleware"]["sample"] == 1
    assert radar["network_os"] == {"key": "network_os", "label": "网络与操作系统", "score": 0.0, "sample": 0}
    assert "technicalCorrectness" not in radar
    assert report["weakestQuestions"][0].startswith(f"[{answer['evaluation']['totalScore']:.1f} 分]")
    assert report["actionItems"]


def test_empty_session_report_has_zero_samples_and_actionable_guidance(tmp_path: Path):
    c = client(tmp_path)
    session_id = c.post(
        "/api/interviews", json={"mode": "KNOWLEDGE", "maxQuestions": 2},
    ).json()["id"]

    response = c.post(f"/api/interviews/{session_id}/report")
    assert response.status_code == 200
    payload = response.json()
    report = payload["report"]
    assert report["overallScore"] == 0
    assert report["answerCount"] == 0
    assert report["weakestQuestions"] == []
    assert report["knowledgeGaps"] == []
    assert report["projectRisks"] == []
    assert all(item["score"] == 0 and item["sample"] == 0 for item in report["abilityRadar"])
    assert any("先完成至少一道题" in item for item in report["actionItems"])
    assert "暂无已评分题目" in payload["markdown"]


def test_replace_retry_and_error_contract(tmp_path: Path):
    c = client(tmp_path)
    resume_id = import_resume(c)["resumeId"]
    session_id = c.post("/api/interviews", json={"resumeId": resume_id, "mode": "PROJECT", "maxQuestions": 3}).json()["id"]
    q1 = c.post(f"/api/interviews/{session_id}/next-question").json()
    q2 = c.post(f"/api/interviews/{session_id}/replace-question").json()
    assert q1["id"] != q2["id"] and q1["sequence"] == q2["sequence"] == 1
    assert q1["content"] != q2["content"]
    answered = c.post(f"/api/interviews/{session_id}/answers", json={"questionId": q2["id"], "content": "我负责核心模块，并通过测试验证。"}).json()
    retried = c.post(f"/api/interviews/{session_id}/retry")
    assert retried.status_code == 200
    assert retried.json()["discardedScore"] == answered["evaluation"]["totalScore"]
    assert c.get(f"/api/interviews/{session_id}/evaluation").status_code == 404
    duplicate_next = c.post(f"/api/interviews/{session_id}/next-question")
    assert duplicate_next.status_code == 409
    assert duplicate_next.json()["code"] == "INVALID_SESSION_STATE"
    replacement = c.post(f"/api/interviews/{session_id}/answers", json={
        "questionId": q2["id"], "content": "首先说明个人职责，然后解释技术机制、边界条件和验证结果。"
    })
    assert replacement.status_code == 200
    c.post(f"/api/interviews/{session_id}/finish", json={"reason": "test"})
    report = c.post(f"/api/interviews/{session_id}/report").json()["report"]
    assert report["answerCount"] == 1


def test_knowledge_search_and_interview_without_resume(tmp_path: Path):
    c = client(tmp_path)
    search = c.get("/api/retrieval/search", params={"q": "volatile 可见性", "mode": "KNOWLEDGE", "topK": 3})
    assert search.status_code == 200
    assert search.json()["chunks"]
    assert all(item["kind"] == "KNOWLEDGE" for item in search.json()["chunks"])
    session = c.post("/api/interviews", json={"mode": "KNOWLEDGE", "maxQuestions": 2})
    assert session.status_code == 200
    assert session.json()["status"] == "INTERVIEW_STARTED"
    question = c.post(f"/api/interviews/{session.json()['id']}/next-question")
    assert question.status_code == 200
    assert question.json()["type"] == "PRINCIPLE"
    state = c.get(f"/api/interviews/{session.json()['id']}").json()
    assert len(state["coveredTopics"]) == 1
    assert "/" in state["coveredTopics"][0]


def test_docx_upload(tmp_path: Path):
    c = client(tmp_path)
    document = Document()
    document.add_heading("项目经历", level=1)
    document.add_paragraph("Python 面试系统")
    document.add_paragraph("使用 FastAPI、Redis 和 Docker 实现服务并编写测试。")
    stream = BytesIO()
    document.save(stream)
    response = c.post("/api/resumes/import", files={
        "file": ("resume.docx", stream.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    })
    assert response.status_code == 200, response.text
    assert response.json()["fileType"] == "docx"
    assert response.json()["factCount"] >= 1


def test_full_interview_advances_stages(tmp_path: Path):
    c = client(tmp_path)
    resume_id = import_resume(c)["resumeId"]
    session_id = c.post("/api/interviews", json={"resumeId": resume_id, "mode": "FULL", "maxQuestions": 4}).json()["id"]
    expected = ["SELF_INTRO", "PROJECT", "PROJECT_DEEP_DIVE", "CANDIDATE_QUESTIONS"]
    for stage in expected:
        question = c.post(f"/api/interviews/{session_id}/next-question").json()
        assert question["stage"] == stage
        result = c.post(f"/api/interviews/{session_id}/answers", json={
            "questionId": question["id"], "content": "首先说明我的个人职责，其次解释技术机制，最后给出测试验证和复盘。"
        })
        assert result.status_code == 200


def test_concurrent_duplicate_answer_is_atomic(tmp_path: Path):
    c = client(tmp_path)
    resume_id = import_resume(c)["resumeId"]
    session_id = c.post("/api/interviews", json={"resumeId": resume_id, "mode": "PROJECT", "maxQuestions": 2}).json()["id"]
    question = c.post(f"/api/interviews/{session_id}/next-question").json()
    engine = c.app.state.engine

    def submit():
        try:
            engine.submit_answer(session_id, question["id"], "首先说明职责，然后解释机制，最后通过测试验证结果。")
            return "ok"
        except Exception as exc:  # the losing request must be the state-machine error
            return getattr(exc, "code", type(exc).__name__)

    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(lambda _: submit(), range(2)))
    assert sorted(results) == ["INVALID_SESSION_STATE", "ok"]
    session = c.get(f"/api/interviews/{session_id}").json()
    assert session["answerCount"] == 1


def test_follow_up_consumes_budget_and_finishes(tmp_path: Path):
    c = client(tmp_path)
    resume_id = import_resume(c)["resumeId"]
    session_id = c.post("/api/interviews", json={"resumeId": resume_id, "mode": "PROJECT", "maxQuestions": 2}).json()["id"]
    question = c.post(f"/api/interviews/{session_id}/next-question").json()
    first = c.post(f"/api/interviews/{session_id}/answers", json={
        "questionId": question["id"], "content": "我主要负责这个模块，但具体机制还需要进一步复习。"
    }).json()
    assert first["nextAction"] == "FOLLOW_UP"
    follow = c.post(f"/api/interviews/{session_id}/follow-up")
    assert follow.status_code == 200, follow.text
    assert follow.json()["sequence"] == 2 and follow.json()["followUp"] is True
    second = c.post(f"/api/interviews/{session_id}/answers", json={
        "questionId": follow.json()["id"], "content": "我会补充关键机制、失败边界以及压测验证方法。"
    }).json()
    assert second["session"]["questionCount"] == 2
    assert second["session"]["followUpCount"] == 1
    assert second["nextAction"] == "FINISH"
    assert second["session"]["status"] == "NEXT_QUESTION"
    exhausted = c.post(f"/api/interviews/{session_id}/next-question")
    assert exhausted.status_code == 409 and "已达到" in exhausted.json()["message"]
    assert c.get(f"/api/interviews/{session_id}").json()["status"] == "FINISHED"


def test_last_answer_can_be_retried_before_budget_is_closed(tmp_path: Path):
    c = client(tmp_path)
    resume_id = import_resume(c)["resumeId"]
    session_id = c.post(
        "/api/interviews",
        json={"resumeId": resume_id, "mode": "PROJECT", "maxQuestions": 1},
    ).json()["id"]
    question = c.post(f"/api/interviews/{session_id}/next-question").json()
    answered = c.post(
        f"/api/interviews/{session_id}/answers",
        json={"questionId": question["id"], "content": "我负责缓存机制，并通过并发测试验证边界。"},
    ).json()
    assert answered["nextAction"] == "FINISH"
    retried = c.post(f"/api/interviews/{session_id}/retry")
    assert retried.status_code == 200
    assert retried.json()["session"]["status"] == "WAITING_ANSWER"
    assert retried.json()["session"]["answerCount"] == 0


def test_answer_length_validation(tmp_path: Path):
    c = client(tmp_path)
    resume_id = import_resume(c)["resumeId"]
    session_id = c.post("/api/interviews", json={"resumeId": resume_id, "mode": "PROJECT"}).json()["id"]
    question = c.post(f"/api/interviews/{session_id}/next-question").json()
    empty = c.post(f"/api/interviews/{session_id}/answers", json={"questionId": question["id"], "content": " "})
    short = c.post(f"/api/interviews/{session_id}/answers", json={"questionId": question["id"], "content": "不会"})
    assert empty.status_code == 400 and empty.json()["code"] == "VALIDATION_ERROR"
    assert short.status_code == 400 and "过短" in short.json()["message"]


def test_invalid_request_body_has_stable_error_contract(tmp_path: Path):
    c = client(tmp_path)
    wrong_shape = c.post("/api/interviews", json=[])
    malformed = c.post("/api/interviews", content=b"{broken", headers={"Content-Type": "application/json"})
    wrong_field_type = c.post("/api/interviews", json={"mode": "PROJECT", "maxQuestions": "abc"})
    wrong_facts_shape = c.put("/api/resumes/missing/facts", json={"facts": {"content": "x"}})
    for response in (wrong_shape, malformed, wrong_field_type, wrong_facts_shape):
        assert response.status_code == 400
        body = response.json()
        assert body["code"] == "INVALID_REQUEST_BODY"
        assert body["path"].startswith("/api/")
        assert body["timestamp"] and isinstance(body["details"], list)


def test_framework_errors_and_upload_limit_use_stable_error_contract(tmp_path: Path):
    c = client(tmp_path)
    missing = c.get("/api/does-not-exist")
    wrong_method = c.patch("/api/health")
    for response, status, code in (
        (missing, 404, "NOT_FOUND"),
        (wrong_method, 405, "METHOD_NOT_ALLOWED"),
    ):
        assert response.status_code == status
        body = response.json()
        assert body["code"] == code and body["path"].startswith("/api/")
        assert body["timestamp"] and body["details"] == []

    settings = replace(
        Settings.load(),
        storage_mode="memory",
        data_dir=tmp_path,
        llm_provider="mock",
        max_file_bytes=16,
    )
    limited = TestClient(create_app(settings))
    oversized = limited.post(
        "/api/resumes/import",
        files={"file": ("large.txt", b"A" * 17, "text/plain")},
    )
    assert oversized.status_code == 413
    assert oversized.json()["code"] == "FILE_TOO_LARGE"


def test_project_mode_never_uses_education_as_project_source(tmp_path: Path):
    c = client(tmp_path)
    source = (Path(__file__).resolve().parents[3] / "java-version" / "src" / "test" / "resources" / "fixtures" /
              "resume-standard.txt").read_text(encoding="utf-8")
    resume_id = c.post("/api/resumes/text", json={"text": source, "fileName": "standard.txt"}).json()["resumeId"]
    detail = c.get(f"/api/resumes/{resume_id}").json()
    project_ids = {fact["id"] for fact in detail["facts"] if fact["type"] == "PROJECT"}
    session_id = c.post("/api/interviews", json={"resumeId": resume_id, "mode": "PROJECT"}).json()["id"]
    question = c.post(f"/api/interviews/{session_id}/next-question").json()
    assert set(question["sourceIds"]).issubset(project_ids)
    assert any(name in question["content"] for name in ("FinanceAgent", "云中摄影"))


def test_heuristic_scoring_is_monotonic_for_poor_and_detailed_answers(tmp_path: Path):
    c = client(tmp_path)
    resume_id = import_resume(c)["resumeId"]

    def score(answer: str) -> float:
        session_id = c.post("/api/interviews", json={"resumeId": resume_id, "mode": "PROJECT", "maxQuestions": 1}).json()["id"]
        question = c.post(f"/api/interviews/{session_id}/next-question").json()
        result = c.post(f"/api/interviews/{session_id}/answers", json={"questionId": question["id"], "content": answer})
        assert result.status_code == 200
        return result.json()["evaluation"]["totalScore"]

    poor = score("这道题我暂时不了解，需要进一步学习。")
    detailed = score("背景是高并发请求。我负责 Redis 缓存和 Kafka 消息链路；首先设计唯一业务键保证幂等，其次用数据库约束兜底，最后通过并发压测、异常回放和监控验证结果，并复盘超时边界。")
    assert detailed > poor
    assert detailed <= 3.5  # no semantic model means heuristic score is deliberately capped


def test_session_lifecycle_and_question_limit_clamping(tmp_path: Path):
    c = client(tmp_path)
    resume_id = import_resume(c)["resumeId"]
    low = c.post("/api/interviews", json={"resumeId": resume_id, "mode": "PROJECT", "maxQuestions": 0}).json()
    high = c.post("/api/interviews", json={"resumeId": resume_id, "mode": "PROJECT", "maxQuestions": 999}).json()
    assert low["maxQuestions"] == 1 and high["maxQuestions"] == 30
    premature = c.post(f"/api/interviews/{low['id']}/finish", json={})
    assert premature.status_code == 200 and premature.json()["status"] == "FINISHED"
    cancelled = c.post(f"/api/interviews/{high['id']}/cancel")
    assert cancelled.status_code == 200 and cancelled.json()["status"] == "CANCELLED"


def test_delete_resume_removes_retrieval_chunks(tmp_path: Path):
    c = client(tmp_path)
    resume_id = import_resume(c)["resumeId"]
    before = c.get("/api/retrieval/search", params={"q": "Redis", "resumeId": resume_id, "mode": "PROJECT"}).json()
    assert before["chunks"]
    deleted = c.delete(f"/api/resumes/{resume_id}")
    assert deleted.status_code == 200 and deleted.json()["ok"] is True
    assert c.get(f"/api/resumes/{resume_id}").status_code == 404
    after = c.get("/api/retrieval/search", params={"q": "Redis", "resumeId": resume_id, "mode": "PROJECT"}).json()
    assert after["empty"] is True and after["chunks"] == []


def test_manual_fact_correction_masks_label_and_metadata(tmp_path: Path):
    c = client(tmp_path)
    resume_id = import_resume(c)["resumeId"]
    updated = c.put(
        f"/api/resumes/{resume_id}/facts",
        json={"facts": [{
            "type": "PROJECT",
            "label": "订单系统 联系人 13812345678",
            "content": "联系邮箱 owner@example.com，负责 Redis 幂等。",
            "metadata": ["主页 https://example.com/private"],
        }]},
    )
    assert updated.status_code == 200
    fact = updated.json()["facts"][0]
    exposed = " ".join([fact["label"], fact["content"], *fact["metadata"]])
    assert "13812345678" not in exposed and "owner@example.com" not in exposed
    assert "example.com/private" not in exposed
