from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from interview_agent.errors import ApiError
from interview_agent.parsing import clean_text, extract_document, extract_facts
from interview_agent.privacy import contains_pii, for_log, mask, scan
from pypdf import PdfWriter


def test_mask_preserves_technical_numbers():
    text = mask("电话 13900001111，HashMap 容量 16，负载因子 0.75，邮箱 a@b.com")
    assert "13900001111" not in text and "a@b.com" not in text
    assert "容量 16" in text and "0.75" in text
    assert not contains_pii(text)


def test_clean_and_extract_sections():
    text, warnings = clean_text("页眉\n页眉\n页眉\n项目经历\n订单中台 2023.01-2023.06\nJava Redis Kafka\n技能\nJava Python")
    assert text.count("页眉") == 0
    facts, confidence, _ = extract_facts(text, "r1")
    assert any(f.type == "PROJECT" for f in facts)
    assert any("Redis" in f.metadata for f in facts)
    assert confidence >= 0.5


def test_headers_are_strict_but_allow_numbering_brackets_and_spaces():
    from interview_agent.parsing import _heading

    assert _heading("我负责技能模块的开发工作。") is None
    assert _heading("二、项目经验") == "PROJECT"
    assert _heading("【专业技能】") == "SKILL"
    assert _heading("技 术 栈") == "SKILL"


def test_separator_and_symbol_noise_are_rejected():
    cleaned, _ = clean_text("张三\n------\n项目经历\n订单系统\n负责接口开发")
    assert "------" not in cleaned
    with pytest.raises(ApiError) as noise:
        clean_text("%%%%%%\n@@@@@@\n------")
    assert noise.value.code == "EMPTY_RESUME_TEXT"


def test_cleaner_normalizes_zero_width_page_counts_and_detects_mojibake():
    cleaned, warnings = clean_text(
        "项目经历\r\n\t技术栈：Java\u200b\r\n1 / 3\r\n锟斤拷锟斤拷锟斤拷\r\n结果：上线"
    )
    assert "\r" not in cleaned and "\t" not in cleaned and "\u200b" not in cleaned
    assert "1 / 3" not in cleaned
    assert "技术栈：Java" in cleaned and "结果：上线" in cleaned
    assert any("乱码" in warning for warning in warnings)


@pytest.mark.parametrize("phone", ["13812345678", "138-1234-5678", "138 1234 5678", "+8613812345678", "010-88886666"])
def test_phone_variants_are_masked(phone: str):
    result = mask(f"电话：{phone}")
    assert phone not in result
    assert "[手机号已隐藏]" in result
    assert not contains_pii(result)


def test_url_address_scan_and_log_safety():
    raw = "地址：北京市海淀区中关村大街 27 号\n主页 https://github.com/example/repo\n手机 13812345678，备用 13900002222，邮箱 a@b.com"
    hits = scan(raw)
    assert hits["phone"] == 2 and hits["email"] == 1 and hits["url"] == 1 and hits["address"] == 1
    safe = mask(raw)
    assert "github.com" not in safe and "中关村" not in safe
    assert "13812345678" not in for_log(raw, 30)
    assert "\n" not in for_log(raw, 30)


FIXTURES = Path(__file__).resolve().parents[3] / "java-version" / "src" / "test" / "resources" / "fixtures"


@pytest.mark.parametrize(
    ("name", "fact_text", "required_type", "tech", "project_count"),
    [
        ("resume-standard.txt", "FinanceAgent", "PROJECT", "pgvector", 2),
        ("resume-tight-headers.txt", "智学助手", "PROJECT", "FastAPI", 2),
        ("resume-numbered-projects.txt", "电商秒杀系统", "PROJECT", "RocketMQ", 2),
        ("resume-minimal.txt", "Java", "SKILL", "Java", 0),
    ],
)
def test_java_regression_fixtures(name: str, fact_text: str, required_type: str, tech: str, project_count: int):
    cleaned, _ = clean_text((FIXTURES / name).read_text(encoding="utf-8"))
    facts, confidence, warnings = extract_facts(cleaned, "fixture")
    combined = "\n".join(f"{fact.label} {fact.content}" for fact in facts)
    stack = [item for fact in facts for item in fact.metadata]
    assert fact_text in combined
    assert any(fact.type == required_type for fact in facts)
    assert tech in stack
    assert sum(fact.type == "PROJECT" for fact in facts) == project_count
    assert len({fact.id for fact in facts}) == len(facts)
    assert not contains_pii(combined)
    if name == "resume-minimal.txt":
        assert confidence < 0.8
        assert any("项目" in item for item in warnings)


def _text_pdf_bytes(text: str) -> bytes:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("ascii")
    objects.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    body = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(body))
        body.extend(f"{index} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(body)
    body.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    body.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        body.extend(f"{offset:010d} 00000 n \n".encode())
    body.extend(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return bytes(body)


def test_document_router_txt_encoding_magic_and_pdf():
    text, kind, _ = extract_document("resume.txt", "项目经历\nPython".encode("gb18030"))
    assert kind == "txt" and "Python" in text
    pdf = _text_pdf_bytes("Resume Python FastAPI")
    text, kind, _ = extract_document("pretend.txt", pdf)
    assert kind == "pdf" and "Python" in text


def test_image_only_pdf_and_unsupported_file_are_explicit():
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    stream = BytesIO()
    writer.write(stream)
    with pytest.raises(ApiError) as image_error:
        extract_document("scan.pdf", stream.getvalue())
    assert image_error.value.code == "UNSUPPORTED_FILE"
    assert "扫描版" in image_error.value.message
    with pytest.raises(ApiError) as type_error:
        extract_document("legacy.doc", b"legacy binary")
    assert type_error.value.code == "UNSUPPORTED_FILE"


def test_password_pdf_empty_and_unknown_binary_are_explicit():
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.encrypt("secret")
    stream = BytesIO()
    writer.write(stream)
    with pytest.raises(ApiError, match="密码"):
        extract_document("protected.pdf", stream.getvalue())
    with pytest.raises(ApiError) as empty:
        extract_document("empty.txt", b"")
    assert empty.value.code == "EMPTY_RESUME_TEXT"
    with pytest.raises(ApiError) as binary:
        extract_document("", bytes([0, 1, 2, 3, 0, 0, 0, 0]))
    assert binary.value.code == "UNSUPPORTED_FILE"


def test_docx_table_text_and_empty_docx():
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "技术栈"
    table.cell(1, 0).text = "订单中台"
    table.cell(1, 1).text = "Java Redis"
    stream = BytesIO()
    document.save(stream)
    text, kind, _ = extract_document("resume.docx", stream.getvalue())
    assert kind == "docx" and "订单中台" in text and "Java Redis" in text

    empty = Document()
    empty_stream = BytesIO()
    empty.save(empty_stream)
    raw, kind, _ = extract_document("empty.docx", empty_stream.getvalue())
    assert kind == "docx" and raw == ""
    with pytest.raises(ApiError):
        clean_text(raw)


def test_tech_stack_long_terms_do_not_shadow_short_terms():
    text, _ = clean_text("专业技能\nJavaScript TypeScript Java 21 Spring Boot Spring Cloud pgvector")
    facts, _, _ = extract_facts(text, "r")
    stack = [item for fact in facts for item in fact.metadata]
    assert all(item in stack for item in ("JavaScript", "TypeScript", "Java", "Spring Boot", "Spring Cloud", "pgvector"))
    only_js, _ = clean_text("专业技能\nJavaScript TypeScript")
    js_facts, _, _ = extract_facts(only_js, "r")
    js_stack = [item for fact in js_facts for item in fact.metadata]
    assert "Java" not in js_stack


def test_fact_extraction_is_deterministic():
    source = (FIXTURES / "resume-standard.txt").read_text(encoding="utf-8")
    cleaned, _ = clean_text(source)
    first = extract_facts(cleaned, "same")
    second = extract_facts(cleaned, "same")
    assert first == second


def test_fact_labels_are_masked_before_they_can_become_citations():
    cleaned, _ = clean_text("项目经历\n订单系统 联系人 13812345678 2024.01-2024.06\n负责 Redis 幂等")
    facts, _, _ = extract_facts(cleaned, "r")
    project = next(fact for fact in facts if fact.type == "PROJECT")
    assert "13812345678" not in project.label
    assert "[手机号已隐藏]" in project.label
